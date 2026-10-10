import os
from datetime import datetime
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QFrame,
    QPushButton, QScrollArea, QComboBox, QFileDialog,
    QMessageBox, QDateEdit, QProgressBar
)
from PyQt6.QtCore import Qt, QDate, QThread, pyqtSignal
from PyQt6.QtGui import QColor
from core.database import get_rapor_data, get_kdv_ozet
from ui.design import (CARD_STYLE, COMBO_STYLE, INPUT_STYLE,
                        TEXT_PRIMARY, TEXT_SECONDARY, TEXT_MUTED,
                        ACCENT_BLUE, ACCENT_GREEN, ACCENT_RED, ACCENT_ORANGE,
                        BG_BASE, BG_SURFACE, BG_INPUT, BORDER,
                        btn_primary, btn_secondary, section_title)


class ExportThread(QThread):
    finished = pyqtSignal(str)
    error    = pyqtSignal(str)

    def __init__(self, export_type, data, kdv, path, period_label):
        super().__init__()
        self.export_type  = export_type
        self.data         = data
        self.kdv          = kdv
        self.path         = path
        self.period_label = period_label

    def run(self):
        try:
            if self.export_type == "pdf":
                self._make_pdf()
            else:
                self._make_excel()
            self.finished.emit(self.path)
        except Exception as e:
            self.error.emit(str(e))

    def _make_pdf(self):
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib import colors
            from reportlab.lib.units import mm
            from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                             Paragraph, Spacer, HRFlowable)
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.enums import TA_LEFT, TA_RIGHT, TA_CENTER
        except ImportError:
            raise Exception("reportlab eksik.\nKomut: pip install reportlab")

        # Türkçe font kayıt — önce uygulama klasörü, sonra sistem
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        import sys, os

        def find_font(names):
            # 1) Uygulamanın kendi assets/fonts klasörü (her OS'ta çalışır)
            base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            for name in names:
                p = os.path.join(base, "assets", "fonts", name)
                if os.path.exists(p):
                    return p
            # 2) Windows sistem fontları
            win = "C:/Windows/Fonts"
            for name in names:
                p = os.path.join(win, name)
                if os.path.exists(p):
                    return p
            # 3) Linux sistem fontları
            linux_dirs = [
                "/usr/share/fonts/truetype/dejavu",
                "/usr/share/fonts/truetype/liberation",
                "/usr/share/fonts/truetype/freefont",
            ]
            for d in linux_dirs:
                for name in names:
                    p = os.path.join(d, name)
                    if os.path.exists(p):
                        return p
            return None

        reg_path  = find_font(["DejaVuSans.ttf", "segoeui.ttf", "arial.ttf", "LiberationSans-Regular.ttf"])
        bold_path = find_font(["DejaVuSans-Bold.ttf", "segoeuib.ttf", "arialbd.ttf", "LiberationSans-Bold.ttf"])

        if reg_path:
            pdfmetrics.registerFont(TTFont("TR", reg_path))
        if bold_path:
            pdfmetrics.registerFont(TTFont("TR-Bold", bold_path))

        FONT_REG  = "TR"      if reg_path  else "Helvetica"
        FONT_BOLD = "TR-Bold" if bold_path else "Helvetica-Bold"

        doc = SimpleDocTemplate(self.path, pagesize=A4,
                                leftMargin=15*mm, rightMargin=15*mm,
                                topMargin=12*mm, bottomMargin=12*mm)
        W = A4[0] - 30*mm
        story = []
        styles = getSampleStyleSheet()

        # Renk tanımları
        DARK    = colors.HexColor("#0f172a")
        BLUE    = colors.HexColor("#3b82f6")
        GREEN   = colors.HexColor("#059669")
        RED     = colors.HexColor("#e04f5f")
        ORANGE  = colors.HexColor("#d9802e")
        GRAY    = colors.HexColor("#64748b")
        LGRAY   = colors.HexColor("#f1f5f9")
        WHITE   = colors.white

        def P(text, font=None, size=10, color="#0f172a", align=TA_LEFT, bold=False):
            f = font or (FONT_BOLD if bold else FONT_REG)
            return Paragraph(
                f'<font name="{f}" size="{size}" color="{color}">{text}</font>',
                ParagraphStyle("p", alignment=align, leading=size*1.4)
            )

        # Başlık bloğu
        header_data = [[
            P("Marjin ERP", font=FONT_BOLD, size=18, color="#ffffff"),
            P(f"Dönem: {self.period_label}<br/>Oluşturulma: {datetime.now().strftime('%d.%m.%Y')}",
              size=9, color="#94a3b8", align=TA_RIGHT)
        ]]
        header_tbl = Table(header_data, colWidths=[W*0.6, W*0.4])
        header_tbl.setStyle(TableStyle([
            ("BACKGROUND",   (0,0), (-1,-1), DARK),
            ("TOPPADDING",   (0,0), (-1,-1), 14),
            ("BOTTOMPADDING",(0,0), (-1,-1), 14),
            ("LEFTPADDING",  (0,0), (0,-1),  14),
            ("RIGHTPADDING", (-1,0),(-1,-1), 14),
            ("VALIGN",       (0,0), (-1,-1), "MIDDLE"),
        ]))
        story.append(header_tbl)
        story.append(HRFlowable(width=W, thickness=3, color=BLUE, spaceAfter=12))

        ozet = self.data["ozet"]

        def sec_label(text):
            return Paragraph(
                f'<font name="{FONT_BOLD}" size="8" color="#94a3b8">{text.upper()}</font>',
                styles["Normal"]
            )

        # Özet metrik kartlar
        story.append(sec_label("Özet Göstergeler"))
        story.append(Spacer(1, 4*mm))

        net_kar = ozet.get("net_kar") or 0
        ciro    = ozet.get("ciro") or 0
        kdv_val = ozet.get("kdv") or 0
        kom     = ozet.get("komisyon") or 0
        siparis = ozet.get("siparis") or 0
        marj    = round(net_kar / ciro * 100, 1) if ciro > 0 else 0

        cw = W / 5
        metrics_data = [
            [P("Net Kar",     size=8, color="#64748b"),
             P("Toplam Ciro", size=8, color="#94a3b8"),
             P("Sipariş",     size=8, color="#94a3b8"),
             P("KDV",         size=8, color="#94a3b8"),
             P("Kar Marjı",   size=8, color="#94a3b8")],
            [P(f"₺{net_kar:,.0f}",  font=FONT_BOLD, size=15, color="#ffffff"),
             P(f"₺{ciro:,.0f}",     font=FONT_BOLD, size=15, color="#0f172a"),
             P(str(siparis),         font=FONT_BOLD, size=15, color="#0f172a"),
             P(f"₺{kdv_val:,.0f}",  font=FONT_BOLD, size=15, color="#0f172a"),
             P(f"%{marj}",           font=FONT_BOLD, size=15, color="#059669")],
        ]
        metrics_tbl = Table(metrics_data, colWidths=[cw]*5)
        metrics_tbl.setStyle(TableStyle([
            ("BACKGROUND",    (0,0), (0,-1), DARK),
            ("BACKGROUND",    (1,0), (-1,-1), LGRAY),
            ("TOPPADDING",    (0,0), (-1,-1), 8),
            ("BOTTOMPADDING", (0,0), (-1,-1), 8),
            ("LEFTPADDING",   (0,0), (-1,-1), 10),
            ("RIGHTPADDING",  (0,0), (-1,-1), 10),
            ("BOX",           (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
            ("INNERGRID",     (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ]))
        story.append(metrics_tbl)
        story.append(Spacer(1, 5*mm))

        # KDV Özeti
        kdv_o = self.kdv
        story.append(sec_label("KDV Özeti"))
        story.append(Spacer(1, 4*mm))

        kdv_rows = [
            [P("Kalem", font=FONT_BOLD, size=9, color="#ffffff"),
             P("Tutar",  font=FONT_BOLD, size=9, color="#ffffff", align=TA_RIGHT)],
            [P("Hesaplanan KDV (Satışlar)",       size=9),
             P(f"₺{kdv_o.get('hesaplanan_kdv',0):,.2f}", size=9, align=TA_RIGHT)],
            [P("— Komisyon KDV'si (indirilecek)", size=9),
             P(f"-₺{kdv_o.get('komisyon_kdv',0):,.2f}", size=9, color="#059669", align=TA_RIGHT)],
            [P("— Kargo KDV'si (indirilecek)",   size=9),
             P(f"-₺{kdv_o.get('kargo_kdv',0):,.2f}", size=9, color="#059669", align=TA_RIGHT)],
            [P("— Platform Hizmet KDV'si",        size=9),
             P(f"-₺{kdv_o.get('platform_kdv',0):,.2f}", size=9, color="#059669", align=TA_RIGHT)],
            [P("— Alış Fatura KDV'si",            size=9),
             P(f"-₺{kdv_o.get('alis_kdv',0):,.2f}", size=9, color="#059669", align=TA_RIGHT)],
            [P("ÖDENECEK NET KDV", font=FONT_BOLD, size=9),
             P(f"₺{kdv_o.get('odenecek_kdv',0):,.2f}", font=FONT_BOLD, size=9, color="#e04f5f", align=TA_RIGHT)],
        ]
        kdv_tbl = Table(kdv_rows, colWidths=[W*0.65, W*0.35])
        kdv_style = TableStyle([
            ("BACKGROUND",    (0,0), (-1,0), DARK),
            ("BACKGROUND",    (0,6), (-1,6), LGRAY),
            ("TOPPADDING",    (0,0), (-1,-1), 6),
            ("BOTTOMPADDING", (0,0), (-1,-1), 6),
            ("LEFTPADDING",   (0,0), (-1,-1), 10),
            ("RIGHTPADDING",  (0,0), (-1,-1), 10),
            ("BOX",           (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
            ("INNERGRID",     (0,0), (-1,-1), 0.5, colors.HexColor("#f1f5f9")),
        ])
        kdv_tbl.setStyle(kdv_style)
        story.append(kdv_tbl)
        story.append(Spacer(1, 5*mm))

        # Ürünler tablosu
        urunler = self.data.get("urunler", [])
        if urunler:
            story.append(sec_label("Ürün Bazlı Performans"))
            story.append(Spacer(1, 4*mm))

            urun_header = [P(h, font=FONT_BOLD, size=8, color="#ffffff")
                           for h in ["Stok Kodu","Ürün Adı","Adet","Ciro","Komisyon","Kargo","Net Kar"]]
            urun_rows = [urun_header]
            for u in urunler[:15]:
                urun_rows.append([
                    P(u.get("stok_kodu",""), size=8),
                    P((u.get("urun_adi","") or "")[:35], size=8),
                    P(str(u.get("toplam_adet",0)), size=8, align=TA_RIGHT),
                    P(f"₺{u.get('ciro',0):,.0f}", size=8, align=TA_RIGHT),
                    P(f"₺{u.get('komisyon',0):,.0f}", size=8, align=TA_RIGHT),
                    P(f"₺{u.get('kargo',0):,.0f}", size=8, align=TA_RIGHT),
                    P(f"₺{u.get('net_kar',0):,.0f}", font=FONT_BOLD, size=8, color="#059669", align=TA_RIGHT),
                ])
            cws = [W*0.1, W*0.32, W*0.06, W*0.13, W*0.12, W*0.11, W*0.12]
            urun_tbl = Table(urun_rows, colWidths=cws)
            urun_tbl.setStyle(TableStyle([
                ("BACKGROUND",    (0,0), (-1,0), DARK),
                ("TOPPADDING",    (0,0), (-1,-1), 5),
                ("BOTTOMPADDING", (0,0), (-1,-1), 5),
                ("LEFTPADDING",   (0,0), (-1,-1), 6),
                ("RIGHTPADDING",  (0,0), (-1,-1), 6),
                ("ROWBACKGROUNDS",(0,1), (-1,-1), [WHITE, LGRAY]),
                ("BOX",           (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
                ("INNERGRID",     (0,0), (-1,-1), 0.5, colors.HexColor("#f1f5f9")),
            ]))
            story.append(urun_tbl)

        # Footer
        story.append(Spacer(1, 8*mm))
        story.append(HRFlowable(width=W, thickness=0.5, color=colors.HexColor("#e2e8f0")))
        story.append(Spacer(1, 2*mm))
        story.append(P(
            f"Marjin ERP · "
            f"Rapor: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
            size=8, color="#94a3b8", align=TA_CENTER
        ))

        doc.build(story)

    def _make_excel(self):
        try:
            import openpyxl
            from openpyxl.styles import (Font, PatternFill, Alignment,
                                          Border, Side, numbers)
        except ImportError:
            raise Exception("openpyxl eksik.")

        wb = openpyxl.Workbook()

        # --- Özet sayfası ---
        ws = wb.active
        ws.title = "Özet"
        ozet = self.data["ozet"]

        dark  = "0F172A"
        blue  = "3B82F6"
        green = "059669"
        red   = "E04F5F"
        lgray = "F1F5F9"

        def hdr(cell, text, bg=dark, fg="FFFFFF", bold=True, sz=11):
            ws[cell] = text
            ws[cell].font = Font(bold=bold, color=fg, size=sz, name="Calibri")
            ws[cell].fill = PatternFill("solid", fgColor=bg)
            ws[cell].alignment = Alignment(horizontal="center", vertical="center")

        def val(cell, text, bold=False, fg="0F172A", align="right"):
            ws[cell] = text
            ws[cell].font = Font(bold=bold, color=fg, size=10, name="Calibri")
            ws[cell].alignment = Alignment(horizontal=align)

        # Başlık
        ws.merge_cells("A1:G1")
        hdr("A1", f"MARJIN ERP — FİNANSAL RAPOR | {self.period_label}", sz=13)
        ws.row_dimensions[1].height = 30

        # Metrikler
        hdr("A3", "ÖZET GÖSTERGELERİ", bg=dark, sz=10)
        ws.merge_cells("A3:G3")
        labels = ["Net Kar", "Toplam Ciro", "Sipariş", "KDV", "Komisyon", "Kargo", "Kar Marjı"]
        ciro = ozet.get("ciro") or 0
        net  = ozet.get("net_kar") or 0
        vals = [
            f"₺{net:,.2f}", f"₺{ciro:,.2f}",
            str(ozet.get("siparis",0)),
            f"₺{(ozet.get('kdv') or 0):,.2f}",
            f"₺{(ozet.get('komisyon') or 0):,.2f}",
            f"₺{(ozet.get('kargo') or 0):,.2f}",
            f"%{round(net/ciro*100,1) if ciro>0 else 0}"
        ]
        for i, (lbl, v) in enumerate(zip(labels, vals)):
            col = chr(65+i)
            hdr(f"{col}4", lbl, bg=lgray, fg="64748B", bold=False, sz=9)
            ws[f"{col}5"] = v
            ws[f"{col}5"].font = Font(bold=True, size=11, name="Calibri",
                                       color=green if i in [0,6] else dark)
            ws[f"{col}5"].alignment = Alignment(horizontal="center")
            ws.column_dimensions[col].width = 16

        # KDV Özeti
        kdv_o = self.kdv
        ws.merge_cells("A7:C7")
        hdr("A7", "KDV ÖZETİ", bg=dark, sz=10)
        kdv_rows = [
            ("Hesaplanan KDV", kdv_o.get("hesaplanan_kdv",0), dark),
            ("Komisyon KDV'si (indirilecek)", kdv_o.get("komisyon_kdv",0), green),
            ("Kargo KDV'si (indirilecek)", kdv_o.get("kargo_kdv",0), green),
            ("Platform Hizmet KDV'si", kdv_o.get("platform_kdv",0), green),
            ("Alış Fatura KDV'si", kdv_o.get("alis_kdv",0), green),
            ("ÖDENECEK NET KDV", kdv_o.get("odenecek_kdv",0), red),
        ]
        for j, (lbl, v, color) in enumerate(kdv_rows):
            r = 8 + j
            ws[f"A{r}"] = lbl
            ws[f"A{r}"].font = Font(bold=(j==5), size=10, name="Calibri")
            ws.merge_cells(f"A{r}:B{r}")
            ws[f"C{r}"] = f"₺{v:,.2f}"
            ws[f"C{r}"].font = Font(bold=(j==5), color=color, size=10, name="Calibri")
            ws[f"C{r}"].alignment = Alignment(horizontal="right")
            if j == 5:
                ws[f"A{r}"].fill = PatternFill("solid", fgColor=lgray)
                ws[f"C{r}"].fill = PatternFill("solid", fgColor=lgray)

        # Ürünler sayfası
        wu = wb.create_sheet("Ürünler")
        urun_headers = ["Stok Kodu","Ürün Adı","Adet","Ciro","Komisyon","Kargo","Net Kar"]
        for i, h in enumerate(urun_headers):
            c = wu.cell(1, i+1, h)
            c.font = Font(bold=True, color="FFFFFF", name="Calibri")
            c.fill = PatternFill("solid", fgColor=dark)
            c.alignment = Alignment(horizontal="center")

        for r, u in enumerate(self.data.get("urunler",[]), 2):
            row = [u.get("stok_kodu",""), u.get("urun_adi",""),
                   u.get("toplam_adet",0), u.get("ciro",0),
                   u.get("komisyon",0), u.get("kargo",0), u.get("net_kar",0)]
            for i, v in enumerate(row):
                c = wu.cell(r, i+1, v)
                c.font = Font(size=10, name="Calibri")
                c.fill = PatternFill("solid", fgColor="FFFFFF" if r%2==0 else lgray)
                if i >= 2:
                    c.alignment = Alignment(horizontal="right")

        for col in wu.columns:
            wu.column_dimensions[col[0].column_letter].width = 18

        wb.save(self.path)


class RaporWidget(QWidget):
    def __init__(self):
        super().__init__()
        self._build_ui()

    def _build_ui(self):
        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.Shape.NoFrame)
        outer.addWidget(scroll)

        content = QWidget()
        content.setStyleSheet(f"background: {BG_BASE};")
        scroll.setWidget(content)

        layout = QVBoxLayout(content)
        layout.setContentsMargins(28, 24, 28, 24)
        layout.setSpacing(20)

        # Header
        header = QHBoxLayout()
        header.addWidget(section_title("Raporlar"))
        layout.addLayout(header)

        # Dönem seçimi
        period_frame = QFrame()
        period_frame.setObjectName("card")
        period_frame.setStyleSheet(CARD_STYLE)
        pfl = QVBoxLayout(period_frame)
        pfl.setContentsMargins(20, 18, 20, 18)
        pfl.setSpacing(14)

        pl_title = QLabel("Dönem Seçimi")
        pl_title.setStyleSheet(f"color: {TEXT_PRIMARY}; font-size: 14px; font-weight: 600;")
        pfl.addWidget(pl_title)

        row1 = QHBoxLayout()
        row1.setSpacing(12)

        self.period_combo = QComboBox()
        self.period_combo.setStyleSheet(COMBO_STYLE)
        self.period_combo.addItem("Tümü", "all")
        self.period_combo.addItem("Bugün", "today")
        self.period_combo.addItem("Bu Hafta", "week")
        self.period_combo.addItem("Bu Ay", "month")
        self.period_combo.addItem("Tarih Aralığı", "custom")
        self.period_combo.currentIndexChanged.connect(self._on_period_change)
        row1.addWidget(QLabel("Dönem:"))
        row1.addWidget(self.period_combo)

        self.bas_edit = QDateEdit()
        self.bas_edit.setStyleSheet(INPUT_STYLE)
        self.bas_edit.setDate(QDate.currentDate().addMonths(-1))
        self.bas_edit.setCalendarPopup(True)
        self.bas_edit.setVisible(False)
        self.bit_edit = QDateEdit()
        self.bit_edit.setStyleSheet(INPUT_STYLE)
        self.bit_edit.setDate(QDate.currentDate())
        self.bit_edit.setCalendarPopup(True)
        self.bit_edit.setVisible(False)

        self.bas_lbl = QLabel("Başlangıç:")
        self.bas_lbl.setStyleSheet(f"color: {TEXT_SECONDARY}; font-size: 12px;")
        self.bas_lbl.setVisible(False)
        self.bit_lbl = QLabel("Bitiş:")
        self.bit_lbl.setStyleSheet(f"color: {TEXT_SECONDARY}; font-size: 12px;")
        self.bit_lbl.setVisible(False)

        row1.addWidget(self.bas_lbl)
        row1.addWidget(self.bas_edit)
        row1.addWidget(self.bit_lbl)
        row1.addWidget(self.bit_edit)
        row1.addStretch()
        pfl.addLayout(row1)

        # Export butonları
        btn_row = QHBoxLayout()
        btn_row.setSpacing(10)

        self.preview_btn = QPushButton("Önizle")
        self.preview_btn.setStyleSheet(btn_secondary())
        self.preview_btn.clicked.connect(self._preview)
        btn_row.addWidget(self.preview_btn)

        pdf_btn = QPushButton("📄  PDF İndir")
        pdf_btn.setStyleSheet(btn_primary())
        pdf_btn.clicked.connect(lambda: self._export("pdf"))
        btn_row.addWidget(pdf_btn)

        excel_btn = QPushButton("📊  Excel İndir")
        excel_btn.setStyleSheet(btn_primary())
        excel_btn.clicked.connect(lambda: self._export("excel"))
        btn_row.addWidget(excel_btn)
        btn_row.addStretch()
        pfl.addLayout(btn_row)
        layout.addWidget(period_frame)

        # Progress
        self.prog_frame = QFrame()
        self.prog_frame.setStyleSheet(f"background: {BG_SURFACE}; border-radius: 6px; border: 1px solid {BORDER};")
        self.prog_frame.setVisible(False)
        prog_l = QVBoxLayout(self.prog_frame)
        prog_l.setContentsMargins(14, 10, 14, 10)
        self.prog_lbl = QLabel("Hazırlanıyor...")
        self.prog_lbl.setStyleSheet(f"color: {TEXT_SECONDARY}; font-size: 12px;")
        self.prog_bar = QProgressBar()
        self.prog_bar.setRange(0, 0)
        self.prog_bar.setFixedHeight(4)
        self.prog_bar.setStyleSheet(f"""
            QProgressBar {{ background: {BG_INPUT}; border-radius: 2px; border: none; }}
            QProgressBar::chunk {{ background: {ACCENT_BLUE}; border-radius: 2px; }}
        """)
        prog_l.addWidget(self.prog_lbl)
        prog_l.addWidget(self.prog_bar)
        layout.addWidget(self.prog_frame)

        # Önizleme kartları
        self.preview_frame = QFrame()
        self.preview_frame.setObjectName("card")
        self.preview_frame.setStyleSheet(CARD_STYLE)
        self.preview_frame.setVisible(False)
        pv_l = QVBoxLayout(self.preview_frame)
        pv_l.setContentsMargins(20, 18, 20, 18)
        pv_l.setSpacing(12)

        pv_title = QLabel("Rapor Önizlemesi")
        pv_title.setStyleSheet(f"color: {TEXT_PRIMARY}; font-size: 14px; font-weight: 600;")
        pv_l.addWidget(pv_title)

        self.pv_cards = QHBoxLayout()
        self.pv_cards.setSpacing(10)
        pv_l.addLayout(self.pv_cards)

        layout.addWidget(self.preview_frame)
        layout.addStretch()

    def _on_period_change(self):
        is_custom = self.period_combo.currentData() == "custom"
        self.bas_lbl.setVisible(is_custom)
        self.bas_edit.setVisible(is_custom)
        self.bit_lbl.setVisible(is_custom)
        self.bit_edit.setVisible(is_custom)

    def _get_params(self):
        period = self.period_combo.currentData()
        bas = self.bas_edit.date().toString("yyyy-MM-dd") if period == "custom" else None
        bit = self.bit_edit.date().toString("yyyy-MM-dd") if period == "custom" else None
        labels = {"all":"Tümü","today":"Bugün","week":"Bu Hafta","month":"Bu Ay","custom":f"{bas} — {bit}"}
        return period, bas, bit, labels.get(period,"Tümü")

    def _preview(self):
        period, bas, bit, label = self._get_params()
        data  = get_rapor_data(period, bas, bit)
        ozet  = data["ozet"]
        ciro  = ozet.get("ciro") or 0
        net   = ozet.get("net_kar") or 0

        # Kartları temizle
        for i in reversed(range(self.pv_cards.count())):
            w = self.pv_cards.itemAt(i).widget()
            if w: w.deleteLater()

        cards = [
            ("Sipariş",    str(ozet.get("siparis",0)),          ACCENT_BLUE),
            ("Ciro",       f"₺{ciro:,.0f}",                     TEXT_PRIMARY),
            ("Net Kar",    f"₺{net:,.0f}",                      ACCENT_GREEN),
            ("KDV",        f"₺{(ozet.get('kdv') or 0):,.0f}",   ACCENT_ORANGE),
            ("Kar Marjı",  f"%{round(net/ciro*100,1) if ciro>0 else 0}", ACCENT_GREEN),
        ]
        for title, value, color in cards:
            frame = QFrame()
            frame.setStyleSheet(f"background: {BG_INPUT}; border-radius: 6px; border: 1px solid {BORDER};")
            fl = QVBoxLayout(frame)
            fl.setContentsMargins(14, 10, 14, 10)
            fl.setSpacing(3)
            t = QLabel(title)
            t.setStyleSheet(f"color: {TEXT_MUTED}; font-size: 10px; font-weight: 600;")
            v = QLabel(value)
            v.setStyleSheet(f"color: {color}; font-size: 18px; font-weight: 600;")
            fl.addWidget(t); fl.addWidget(v)
            self.pv_cards.addWidget(frame)

        self.preview_frame.setVisible(True)

    def _export(self, export_type):
        period, bas, bit, label = self._get_params()

        ext   = "pdf" if export_type == "pdf" else "xlsx"
        fname = f"MarjinERP_Rapor_{datetime.now().strftime('%Y%m%d_%H%M')}.{ext}"
        ftype = "PDF Dosyası (*.pdf)" if ext == "pdf" else "Excel Dosyası (*.xlsx)"

        path, _ = QFileDialog.getSaveFileName(self, "Raporu Kaydet", fname, ftype)
        if not path:
            return

        data = get_rapor_data(period, bas, bit)
        kdv  = get_kdv_ozet()

        self.prog_frame.setVisible(True)
        self.prog_lbl.setText(f"{ext.upper()} oluşturuluyor...")

        self.thread = ExportThread(export_type, data, kdv, path, label)
        self.thread.finished.connect(self._on_done)
        self.thread.error.connect(self._on_error)
        self.thread.start()

    def _on_done(self, path):
        self.prog_frame.setVisible(False)
        reply = QMessageBox.information(self, "Hazır",
            f"Rapor oluşturuldu:\n{path}\n\nAçmak ister misiniz?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
        if reply == QMessageBox.StandardButton.Yes:
            import subprocess, sys
            if sys.platform == "win32":
                os.startfile(path)
            else:
                subprocess.Popen(["xdg-open", path])

    def _on_error(self, msg):
        self.prog_frame.setVisible(False)
        QMessageBox.critical(self, "Hata", f"Rapor oluşturulamadı:\n\n{msg}")
